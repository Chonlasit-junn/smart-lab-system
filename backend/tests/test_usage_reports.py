import asyncio
import io
import os
import sys
import unittest
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

import httpx
from fastapi import FastAPI
from jose import jwt
from openpyxl import load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES
from sqlalchemy import BigInteger, create_engine, event
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool


BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))
os.environ["SQLALCHEMY_DATABASE_URL"] = "sqlite:///:memory:"

import models  # noqa: E402
from database import Base, get_db  # noqa: E402
from routers import reports, users  # noqa: E402
from usage_report import LABELS, ReportFilters, build_usage_workbook  # noqa: E402


@compiles(JSONB, "sqlite")
def compile_jsonb(_type, _compiler, **_kwargs):
    return "TEXT"


@compiles(BigInteger, "sqlite")
def compile_bigint(_type, _compiler, **_kwargs):
    return "INTEGER"


class UsageReportTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()
        self.now = datetime(2026, 10, 25, 5, tzinfo=timezone.utc)
        self.start = datetime(2026, 10, 23, 16, 59, tzinfo=timezone.utc)  # 23:59 Bangkok
        self.admin = models.User(first_name="Admin", last_name="Test", email="admin@example.test", password="not-exported")
        self.user = models.User(first_name="=1+1", last_name="ผู้ใช้", email="student@example.test", password="not-exported")
        self.other_user = models.User(first_name="Other", last_name="Person", email="other@example.test", password="not-exported")
        self.lab = models.Lab(code="LAB01", name="ห้อง Lab", status="active")
        self.other_lab = models.Lab(code="LAB02", name="Other lab", status="active")
        self.admin_role = models.Role(name="admin")
        self.db.add_all([self.admin, self.user, self.other_user, self.lab, self.other_lab, self.admin_role])
        self.db.flush()
        self.admin.roles = [self.admin_role]
        self.device = models.LabDevice(device_id="lab-device-1", device_name="Renamed device", lab_id=self.lab.id, agent_token_hash="never-export-me")
        self.other_device = models.LabDevice(device_id="lab-device-2", device_name="Other device", lab_id=self.other_lab.id, agent_token_hash="other-private-hash")
        self.db.add_all([self.device, self.other_device])
        self.db.flush()
        self.session = models.LabAccessLog(user_id=self.user.id, lab_id=self.lab.id, lab_device_id=self.device.id,
            device_used="Original workstation", access_type="manual", status="success", session_status="completed",
            entry_time=self.start, exit_time=self.start + timedelta(minutes=3), end_reason="logout")
        self.other_session = models.LabAccessLog(user_id=self.other_user.id, lab_id=self.other_lab.id, lab_device_id=self.other_device.id,
            device_used="Other workstation", access_type="manual", status="success", session_status="completed",
            entry_time=self.start, exit_time=self.start + timedelta(minutes=2))
        self.db.add_all([self.session, self.other_session])
        self.db.flush()
        self.db.add_all([
            models.Student(user_id=self.user.id, student_id="00123456"),
            models.UserPassport(user_id=self.user.id, phone="0812345678"),
            models.ProgramUsageLog(lab_access_log_id=self.session.id, program_name="notepad.exe", usage_start_time=self.start,
                usage_end_time=self.start + timedelta(minutes=2), duration_seconds=120, event_id="usage-1", created_at=self.now),
            models.ProgramUsageLog(lab_access_log_id=self.session.id, program_name="notepad.exe", usage_start_time=self.start + timedelta(minutes=2),
                usage_end_time=self.start + timedelta(minutes=3), duration_seconds=60, event_id="usage-2", created_at=self.now),
            models.ProgramUsageLog(lab_access_log_id=self.other_session.id, program_name="chrome.exe", usage_start_time=self.start,
                usage_end_time=self.start + timedelta(minutes=2), duration_seconds=120, event_id="usage-3"),
            models.UsageViolation(lab_access_log_id=self.session.id, program_name="blocked.example", reason="=HYPERLINK(\"https://example.test\")",
                action_taken="logout", detection_source="browser_domain", detected_at=self.start + timedelta(minutes=1), event_id="violation-1"),
            models.UsageViolation(lab_access_log_id=self.session.id, program_name="blocked.exe", reason="Raw reason",
                action_taken="logout", detection_source="executable_identity", detected_at=self.start + timedelta(minutes=2), event_id="violation-2"),
        ])
        self.db.commit()
        self.app = FastAPI()
        self.app.include_router(reports.router)
        self.app.dependency_overrides[get_db] = lambda: self.db
        self.statements = []
        event.listen(self.engine, "before_cursor_execute", self._record_sql)

    def _record_sql(self, _connection, _cursor, statement, _parameters, _context, _many):
        self.statements.append(statement)

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def filters(self, **changes):
        values = dict(date_from=date(2026, 10, 23), date_to=date(2026, 10, 24), language="en")
        values.update(changes)
        return ReportFilters(**values)

    def workbook(self, **changes):
        handle = build_usage_workbook(self.db, self.filters(**changes), now=self.now)
        try:
            return load_workbook(io.BytesIO(handle.read()))
        finally:
            handle.close()

    def request(self, path, user="admin"):
        headers = {}
        if user:
            account = self.admin if user == "admin" else self.user
            token = jwt.encode({"sub": account.email, "auth_version": account.auth_version}, users.SECRET_KEY, algorithm=users.ALGORITHM)
            headers["Authorization"] = f"Bearer {token}"

        async def send():
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app), base_url="http://test") as client:
                return await client.get(path, headers=headers)

        with patch("usage_report.utc_now", return_value=self.now):
            return asyncio.run(send())

    def test_four_sheets_midnight_split_and_checkpoint_aggregation(self):
        book = self.workbook(user_id=self.user.id)
        self.assertEqual(book.sheetnames, list(LABELS["en"]["sheets"]))
        rows = list(book["Program usage"].values)
        self.assertEqual(len(rows), 3)
        self.assertEqual([row[10].date() for row in rows[1:]], [date(2026, 10, 23), date(2026, 10, 24)])
        self.assertEqual([row[14] for row in rows[1:]], [60, 120])
        self.assertEqual(rows[2][15], timedelta(seconds=120))
        self.assertEqual(book["Summary"]["B16"].value, 180)
        self.assertEqual(book["Summary"]["B17"].value, 2)
        self.assertEqual(len(list(book["Violations"].values)), 3)  # two violations do not multiply usage
        self.assertTrue(book["Program usage"].freeze_panes)
        self.assertEqual(book["Program usage"].auto_filter.ref, "A1:P3")
        book.close()

    def test_snapshot_identity_text_and_typed_dates_are_preserved(self):
        book = self.workbook(user_id=self.user.id)
        sheet = book["Program usage"]
        self.assertEqual(sheet["C2"].value, "=1+1 ผู้ใช้")
        self.assertEqual(sheet["C2"].data_type, "s")
        self.assertEqual(sheet["E2"].value, "00123456")
        self.assertEqual(sheet["F2"].value, "0812345678")
        self.assertEqual(sheet["J2"].value, "Original workstation")
        self.assertEqual(sheet["K2"].number_format, "dd/mm/yyyy")
        self.assertEqual(sheet["M2"].value, datetime(2026, 10, 23, 23, 59))
        for sheet in book:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
                    self.assertNotIn("never-export-me", str(cell.value))
                    self.assertNotIn("not-exported", str(cell.value))
        book.close()

    def test_each_filter_and_combinations_apply_to_all_sheets(self):
        for changes in ({"lab_id": self.lab.id}, {"device_id": self.device.id}, {"user_id": self.user.id},
                        {"lab_id": self.lab.id, "device_id": self.device.id, "user_id": self.user.id}):
            with self.subTest(changes=changes):
                book = self.workbook(**changes)
                self.assertEqual(book["Sessions"].max_row, 2)
                self.assertEqual(book["Program usage"].max_row, 3)
                self.assertEqual(book["Violations"].max_row, 3)
                self.assertEqual(book["Summary"]["B12"].value, 1)
                book.close()
        book = self.workbook(lab_id=self.other_lab.id, user_id=self.user.id)
        self.assertEqual(book["Sessions"].max_row, 1)
        book.close()

    def test_only_requested_bangkok_day_and_actual_event_times_not_receipt(self):
        book = self.workbook(date_from=date(2026, 10, 24), date_to=date(2026, 10, 24), user_id=self.user.id)
        self.assertEqual(book["Program usage"]["O2"].value, 120)
        self.assertEqual(book["Summary"]["B15"].value, 120)
        self.assertEqual(book["Violations"]["L2"].value, datetime(2026, 10, 24, 0, 0))
        book.close()

    def test_empty_report_retains_four_sheets_headers_and_zero_totals(self):
        book = self.workbook(date_from=date(2026, 1, 1), date_to=date(2026, 1, 1), language="th")
        self.assertEqual(book.sheetnames, list(LABELS["th"]["sheets"]))
        for sheet in list(book)[1:]:
            self.assertEqual(sheet.max_row, 1)
        self.assertEqual(book["สรุป"]["B12"].value, 0)
        self.assertEqual(book["สรุป"]["B16"].value, 0)
        book.close()

    def test_export_is_read_only_and_does_not_autoflush(self):
        self.user.first_name = "Pending change"
        self.statements.clear()
        book = self.workbook()
        book.close()
        self.assertFalse(any(sql.lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE", "ALTER", "CREATE")) for sql in self.statements))
        self.assertIn(self.user, self.db.dirty)
        self.assertLessEqual(len(self.statements), 3)

    def test_export_more_than_100_rows_without_ui_pagination_or_n_plus_one(self):
        self.db.add_all([models.ProgramUsageLog(lab_access_log_id=self.session.id, program_name=f"program-{i}.exe",
            usage_start_time=self.start, usage_end_time=self.start + timedelta(seconds=1), duration_seconds=1,
            event_id=f"many-{i}") for i in range(125)])
        self.db.commit()
        self.statements.clear()
        book = self.workbook()
        self.assertGreater(book["Program usage"].max_row, 125)
        self.assertEqual(len(self.statements), 3)
        book.close()

    def test_active_sessions_clipped_at_generation_time_and_gatekeeper_excluded(self):
        self.session.entry_time = self.now - timedelta(minutes=5)
        self.session.exit_time = None
        self.session.session_status = "active"
        self.other_session.access_type = "entry"
        self.db.commit()
        book = self.workbook(date_from=date(2026, 10, 25), date_to=date(2026, 10, 25))
        self.assertEqual(book["Sessions"]["O2"].value, 300)
        self.assertIsNone(book["Sessions"]["L2"].value)
        self.assertEqual(book["Summary"]["B14"].value, 1)
        book.close()

    def test_invalid_range_language_ids_and_source_limit_are_explicit_errors(self):
        for changes in ({"date_to": date(2026, 10, 22)}, {"date_to": date(2026, 11, 23)}, {"language": "xx"}, {"lab_id": 0}):
            with self.subTest(changes=changes), self.assertRaises(Exception) as raised:
                self.filters(**changes)
            self.assertEqual(raised.exception.status_code, 422)
        before = set(ALL_TEMP_FILES)
        with patch("usage_report.MAX_SOURCE_ROWS", 1), self.assertRaises(Exception) as raised:
            self.workbook()
        self.assertEqual(raised.exception.status_code, 413)
        self.assertEqual(set(ALL_TEMP_FILES), before)

    def test_route_authentication_and_role_checks_for_every_endpoint(self):
        for endpoint in ("options", "users?q=student", "usage.xlsx?date_from=2026-10-23&date_to=2026-10-24"):
            self.assertEqual(self.request(f"/admin/reports/{endpoint}", user=None).status_code, 401)
            self.assertEqual(self.request(f"/admin/reports/{endpoint}", user="student").status_code, 403)

    def test_download_headers_and_real_xlsx_response(self):
        self.statements.clear()
        response = self.request("/admin/reports/usage.xlsx?date_from=2026-10-23&date_to=2026-10-24&language=en")
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else "")
        self.assertIn("spreadsheetml", response.headers["content-type"])
        self.assertIn("2026-10-23_2026-10-24_en.xlsx", response.headers["content-disposition"])
        self.assertEqual(response.headers["cache-control"], "no-store")
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(len(book.sheetnames), 4)
        book.close()
        self.assertFalse(any(sql.lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE")) for sql in self.statements))

    def test_options_and_bounded_search_do_not_expose_credentials(self):
        response = self.request("/admin/reports/options")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["devices"]), 2)
        self.assertNotIn("agent_token_hash", response.text)
        self.assertEqual(self.request("/admin/reports/users?q=s").json()["data"], [])
        self.assertEqual(self.request("/admin/reports/users?q=student").json()["data"][0]["id"], self.user.id)
        self.assertEqual(self.request("/admin/reports/users?q=%25%25").json()["data"], [])

    def test_route_validation_missing_filter_and_concurrency_limit(self):
        base = "/admin/reports/usage.xlsx?date_from=2026-10-23&date_to=2026-10-24"
        for query in ("&language=xx", "&lab_id=0", "&user_id=abc"):
            self.assertEqual(self.request(base + query).status_code, 422)
        self.assertEqual(self.request(base + "&lab_id=999999").status_code, 404)
        with patch.object(reports, "_export_slots") as semaphore:
            semaphore.acquire.return_value = False
            response = self.request(base)
            self.assertEqual(response.status_code, 429)
            self.assertEqual(response.headers["retry-after"], "10")
            semaphore.release.assert_not_called()

    def test_duplicate_contact_rows_do_not_multiply_activity_and_search_is_bounded(self):
        self.db.add_all([
            models.Student(user_id=self.user.id, student_id="99123456"),
            models.UserPassport(user_id=self.user.id, phone="0999999999"),
            *[models.User(first_name="Matching", last_name=str(i), email=f"matching-{i}@example.test", password="private") for i in range(25)],
        ])
        self.db.commit()
        book = self.workbook(user_id=self.user.id)
        self.assertEqual(book["Program usage"].max_row, 3)
        self.assertEqual(book["Sessions"].max_row, 2)
        self.assertEqual(book["Summary"]["B16"].value, 180)
        book.close()
        data = self.request("/admin/reports/users?q=matching").json()
        self.assertEqual(len(data["data"]), 20)
        self.assertTrue(data["has_more"])

    def test_warning_count_and_database_null_device_remain_explicit(self):
        self.session.lab_device_id = None
        self.session.exit_time = self.start - timedelta(minutes=1)
        self.db.commit()
        book = self.workbook(user_id=self.user.id)
        self.assertEqual(book["Summary"]["B13"].value, 0)
        self.assertEqual(book["Summary"]["B18"].value, 1)
        self.assertEqual(book["Summary"]["B16"].value, 180)
        self.assertIsNone(book["Sessions"]["I2"].value)
        book.close()

    def test_stream_cancellation_closes_spooled_file(self):
        handle = io.BytesIO(b"x" * 100_000)
        stream = reports._file_chunks(handle)
        self.assertEqual(len(next(stream)), 64 * 1024)
        stream.close()
        self.assertTrue(handle.closed)

    def test_zero_length_session_still_counts_its_user_and_device_without_inventing_time(self):
        self.session.entry_time = self.now - timedelta(minutes=1)
        self.session.exit_time = self.session.entry_time
        self.db.commit()
        book = self.workbook(date_from=date(2026, 10, 25), date_to=date(2026, 10, 25), user_id=self.user.id)
        for address in ("B12", "B13", "B14"):
            self.assertEqual(book["Summary"][address].value, 1)
        self.assertEqual(book["Summary"]["B15"].value, 0)
        self.assertEqual(book["Summary"]["B21"].value, 1)
        book.close()


if __name__ == "__main__":
    unittest.main()
