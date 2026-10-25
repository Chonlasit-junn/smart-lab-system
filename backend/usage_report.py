"""Read-only Agent reporting shared by downloadable report formats."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from tempfile import SpooledTemporaryFile

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, func, or_

import models
from time_utils import LAB_TIMEZONE, utc_now


MAX_REPORT_DAYS = 31
MAX_SOURCE_ROWS = 50_000
BATCH_SIZE = 500
DATE_FORMAT = "dd/mm/yyyy"
DATETIME_FORMAT = "dd/mm/yyyy hh:mm:ss"
LABELS = {
    "en": {
        "sheets": ("Summary", "Program usage", "Sessions", "Violations"),
        "common": ("Session ID", "User ID", "Name", "Email", "Student ID", "Phone",
                   "Lab code", "Lab name", "Device ID", "Device name"),
        "usage": ("Date", "Program", "First observed", "Last observed", "Seconds", "Duration"),
        "session": ("Entry time", "Exit time", "Status", "End reason", "Seconds in period", "Duration in period"),
        "violation": ("Violation ID", "Detected at", "Program / domain", "Reason", "Action", "Detection source"),
        "daily": ("Date", "Users", "Registered devices", "Sessions", "Session seconds", "Program seconds", "Violations"),
        "title": "Smart Lab usage report", "from": "From", "to": "Through",
        "generated": "Generated at", "timezone": "Timezone", "lab": "Lab ID filter",
        "device": "Device ID filter", "user": "User ID filter", "all": "All",
        "users": "Distinct users", "devices": "Distinct registered devices",
        "sessions": "Distinct sessions", "session_seconds": "Session seconds in period",
        "usage_seconds": "Foreground program seconds", "violations": "Recorded violations",
        "invalid": "Invalid time intervals excluded", "notes": "Notes",
        "definition": "Agent sessions only. Program time measures foreground windows, not occupancy.",
        "snapshot": "Snapshot at generation time; delayed Agent deliveries may change a later export.",
        "history": "Device names use session snapshots. Names and contacts reflect current account data.",
    },
    "th": {
        "sheets": ("สรุป", "การใช้โปรแกรม", "Session", "การผิดกฎ"),
        "common": ("Session ID", "User ID", "ชื่อผู้ใช้", "อีเมล", "รหัสนักศึกษา", "เบอร์โทรศัพท์",
                   "รหัส Lab", "ชื่อ Lab", "Device ID", "ชื่อเครื่อง"),
        "usage": ("วันที่", "โปรแกรม", "เริ่มใช้งาน", "สิ้นสุดใช้งาน", "วินาที", "ระยะเวลา"),
        "session": ("เวลาเข้า", "เวลาออก", "สถานะ", "เหตุผลที่จบ", "วินาทีในช่วงรายงาน", "ระยะเวลาในช่วงรายงาน"),
        "violation": ("Violation ID", "เวลาตรวจพบ", "โปรแกรม / โดเมน", "เหตุผล", "การดำเนินการ", "แหล่งตรวจพบ"),
        "daily": ("วันที่", "ผู้ใช้", "เครื่องที่ลงทะเบียน", "Session", "วินาทีใน Session", "วินาทีใช้โปรแกรม", "การผิดกฎ"),
        "title": "รายงานการใช้งาน Smart Lab", "from": "ตั้งแต่วันที่", "to": "ถึงวันที่",
        "generated": "สร้างรายงานเมื่อ", "timezone": "เขตเวลา", "lab": "ตัวกรอง Lab ID",
        "device": "ตัวกรอง Device ID", "user": "ตัวกรอง User ID", "all": "ทั้งหมด",
        "users": "ผู้ใช้ไม่ซ้ำ", "devices": "เครื่องที่ลงทะเบียนไม่ซ้ำ",
        "sessions": "Session ไม่ซ้ำ", "session_seconds": "วินาทีใน Session ภายในช่วงรายงาน",
        "usage_seconds": "วินาทีใช้โปรแกรมด้านหน้า", "violations": "การผิดกฎที่บันทึกแล้ว",
        "invalid": "ช่วงเวลาผิดปกติที่ไม่นำมารวม", "notes": "หมายเหตุ",
        "definition": "เฉพาะ Session ของ Agent เวลาโปรแกรมคือเวลาหน้าต่างด้านหน้า ไม่ใช่เวลาครองเครื่อง",
        "snapshot": "ข้อมูล ณ เวลาสร้างรายงาน ข้อมูลที่ Agent ส่งย้อนหลังอาจทำให้รายงานครั้งถัดไปเปลี่ยนแปลง",
        "history": "ชื่อเครื่องใช้ค่าที่บันทึกใน Session ส่วนชื่อและข้อมูลติดต่อเป็นข้อมูลบัญชีปัจจุบัน",
    },
}


@dataclass(frozen=True)
class ReportFilters:
    date_from: date
    date_to: date
    lab_id: int | None = None
    device_id: int | None = None
    user_id: int | None = None
    language: str = "th"

    def __post_init__(self):
        if self.date_to < self.date_from or (self.date_to - self.date_from).days >= MAX_REPORT_DAYS:
            raise HTTPException(422, detail="REPORT_DATE_RANGE")
        if self.date_to == date.max or self.language not in LABELS:
            raise HTTPException(422, detail="REPORT_FILTER_INVALID")
        if any(value is not None and (not isinstance(value, int) or value < 1)
               for value in (self.lab_id, self.device_id, self.user_id)):
            raise HTTPException(422, detail="REPORT_FILTER_INVALID")

    @property
    def start(self):
        return datetime.combine(self.date_from, time.min, LAB_TIMEZONE).astimezone(timezone.utc)

    @property
    def end(self):
        return datetime.combine(self.date_to + timedelta(days=1), time.min, LAB_TIMEZONE).astimezone(timezone.utc)


def _utc(value):
    # These are database timestamptz values; SQLite tests return naive UTC.
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def _excel_time(value):
    return _utc(value).astimezone(LAB_TIMEZONE).replace(tzinfo=None) if value else None


def _parts(start, end, seconds, lower, upper):
    """Allocate integer seconds across Bangkok midnights without rounding drift."""
    start, end = _utc(start), _utc(end)
    if end <= start or seconds < 0:
        return
    left, stop = max(start, lower), min(end, upper)
    total = (end - start).total_seconds()
    while left < stop:
        local_day = left.astimezone(LAB_TIMEZONE).date()
        midnight = datetime.combine(local_day + timedelta(days=1), time.min, LAB_TIMEZONE)
        right = min(stop, midnight.astimezone(timezone.utc))
        before = round(seconds * (left - start).total_seconds() / total)
        after = round(seconds * (right - start).total_seconds() / total)
        yield local_day, left, right, after - before
        left = right


def _session_query(db, filters, until):
    access = models.LabAccessLog
    # Aggregate contacts once, not a correlated profile scan for every session.
    # Legacy duplicate profile rows must not multiply usage/session records.
    student = db.query(models.Student.user_id, func.min(models.Student.student_id).label("student_id")).group_by(models.Student.user_id).subquery()
    phone = db.query(models.UserPassport.user_id, func.min(models.UserPassport.phone).label("phone")).group_by(models.UserPassport.user_id).subquery()
    usage_exists = db.query(models.ProgramUsageLog.id).filter(
        models.ProgramUsageLog.lab_access_log_id == access.id,
        models.ProgramUsageLog.usage_start_time < until,
        models.ProgramUsageLog.usage_end_time > filters.start,
    ).exists()
    violation_exists = db.query(models.UsageViolation.id).filter(
        models.UsageViolation.lab_access_log_id == access.id,
        models.UsageViolation.detected_at >= filters.start,
        models.UsageViolation.detected_at < until,
    ).exists()
    query = db.query(
        access.id.label("session_id"), access.user_id, access.lab_device_id,
        access.device_used, access.entry_time, access.exit_time, access.session_status, access.end_reason,
        models.User.first_name, models.User.last_name, models.User.email,
        student.c.student_id, phone.c.phone,
        models.Lab.code.label("lab_code"), models.Lab.name.label("lab_name"),
        models.LabDevice.device_id.label("device_identifier"),
    ).join(models.User, models.User.id == access.user_id).join(
        models.Lab, models.Lab.id == access.lab_id,
    ).outerjoin(models.LabDevice, models.LabDevice.id == access.lab_device_id).outerjoin(
        student, student.c.user_id == access.user_id,
    ).outerjoin(phone, phone.c.user_id == access.user_id).filter(
        access.access_type == "manual", access.status == "success",
        or_(
            and_(access.entry_time < until, or_(access.exit_time > filters.start, access.exit_time.is_(None))),
            usage_exists, violation_exists,
        ),
    )
    for column, value in ((access.lab_id, filters.lab_id), (access.lab_device_id, filters.device_id), (access.user_id, filters.user_id)):
        if value is not None:
            query = query.filter(column == value)
    return query


def _common(row):
    return [row.session_id, row.user_id, f"{row.first_name} {row.last_name}", row.email,
            row.student_id, row.phone, row.lab_code, row.lab_name,
            row.device_identifier, row.device_used]


def _append(sheet, values, *, header=False):
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet)
        if isinstance(value, str):
            # Literal strings, even '=...', must never become executable formulas.
            cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
            cell.data_type = "s"
            cell.number_format = "@"
        else:
            cell.value = value
            if isinstance(value, datetime):
                cell.number_format = DATETIME_FORMAT
            elif isinstance(value, date):
                cell.number_format = DATE_FORMAT
            elif isinstance(value, timedelta):
                cell.number_format = "[h]:mm:ss"
            elif isinstance(value, (int, float)):
                cell.number_format = "#,##0"
        cell.font = Font(name="Arial", size=11, bold=header, color="FFFFFF" if header else "17202A")
        cell.alignment = Alignment(vertical="center", horizontal="center" if header else None, wrap_text=True)
        if header:
            cell.fill = PatternFill("solid", fgColor="26364A")
        cells.append(cell)
    sheet.append(cells)


def build_usage_workbook(db, filters, *, now=None):
    """Generate a bounded snapshot using SELECTs only; caller owns returned file."""
    now = _utc(now or utc_now())
    until = min(filters.end, now)
    labels = LABELS[filters.language]
    workbook = Workbook(write_only=True)
    sheets = [workbook.create_sheet(name) for name in labels["sheets"]]
    summary, usage_sheet, session_sheet, violation_sheet = sheets
    for sheet in sheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_format.defaultRowHeight = 30
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.tabColor = "26364A" if sheet == summary else "64748B"
    headers = [labels["common"] + labels[kind] for kind in ("usage", "session", "violation")]
    for sheet, heading in zip(sheets[1:], headers):
        sheet.freeze_panes = "D2"
        sheet.print_title_rows = "1:1"
        for index, name in enumerate(heading, 1):
            sheet.column_dimensions[get_column_letter(index)].width = 25 if index > 10 else (34 if index in (3, 4, 8, 10) else 20)
        sheet.row_dimensions[1].height = 34
        _append(sheet, heading, header=True)
    source_rows = 0
    invalid = 0

    def bounded(query):
        nonlocal source_rows
        for row in query.limit(MAX_SOURCE_ROWS + 1).yield_per(BATCH_SIZE):
            source_rows += 1
            if source_rows > MAX_SOURCE_ROWS:
                raise HTTPException(413, detail="REPORT_TOO_LARGE")
            yield row

    days = {}
    cursor = filters.date_from
    while cursor <= filters.date_to:
        days[cursor] = {"users": set(), "devices": set(), "sessions": set(), "session_seconds": 0, "usage_seconds": 0, "violations": 0}
        cursor += timedelta(days=1)

    def mark(day, row):
        target = days[day]
        target["users"].add(row.user_id)
        target["sessions"].add(row.session_id)
        if row.lab_device_id is not None:
            target["devices"].add(row.lab_device_id)
        return target

    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    try:
        with db.no_autoflush:
            for model, identifier in ((models.Lab, filters.lab_id), (models.LabDevice, filters.device_id), (models.User, filters.user_id)):
                if identifier is not None and db.query(model.id).filter(model.id == identifier).first() is None:
                    raise HTTPException(404, detail="REPORT_FILTER_NOT_FOUND")
            session_query = _session_query(db, filters, until)
            sessions = {}
            for row in bounded(session_query.order_by(models.LabAccessLog.id)):
                sessions[row.session_id] = row
                seconds = 0
                if row.entry_time:
                    start, end = _utc(row.entry_time), _utc(row.exit_time) if row.exit_time else now
                    if filters.start <= start < until:
                        mark(start.astimezone(LAB_TIMEZONE).date(), row)
                    if end < start:
                        invalid += 1
                    else:
                        for day, _, _, part in _parts(start, end, int((end - start).total_seconds()), filters.start, until):
                            mark(day, row)["session_seconds"] += part
                            seconds += part
                else:
                    invalid += 1
                _append(session_sheet, _common(row) + [_excel_time(row.entry_time), _excel_time(row.exit_time),
                        row.session_status, row.end_reason, seconds, timedelta(seconds=seconds)])
            # The scoped subquery avoids a giant IN list and shares all filters.
            scoped_ids = session_query.with_entities(models.LabAccessLog.id).subquery()
            usage_query = db.query(models.ProgramUsageLog).join(
                scoped_ids, scoped_ids.c.id == models.ProgramUsageLog.lab_access_log_id,
            ).filter(models.ProgramUsageLog.usage_start_time < until, models.ProgramUsageLog.usage_end_time > filters.start)
            aggregates = {}
            for log in bounded(usage_query.order_by(models.ProgramUsageLog.id)):
                row = sessions.get(log.lab_access_log_id)
                if row is None:
                    continue  # A concurrent new session is not part of this snapshot.
                if not log.usage_start_time or not log.usage_end_time or _utc(log.usage_end_time) <= _utc(log.usage_start_time) or log.duration_seconds < 0:
                    invalid += 1
                    continue
                for day, left, right, seconds in _parts(log.usage_start_time, log.usage_end_time, log.duration_seconds, filters.start, until):
                    mark(day, row)["usage_seconds"] += seconds
                    key = (day, row.session_id, log.program_name)
                    if key not in aggregates:
                        if len(aggregates) >= MAX_SOURCE_ROWS:
                            raise HTTPException(413, detail="REPORT_TOO_LARGE")
                        aggregates[key] = [left, right, 0]
                    aggregate = aggregates[key]
                    aggregate[0], aggregate[1] = min(aggregate[0], left), max(aggregate[1], right)
                    aggregate[2] += seconds
            for (day, session_id, program), (start, end, seconds) in sorted(aggregates.items()):
                _append(usage_sheet, _common(sessions[session_id]) + [day, program, _excel_time(start), _excel_time(end), seconds, timedelta(seconds=seconds)])
            violations = db.query(models.UsageViolation).join(
                scoped_ids, scoped_ids.c.id == models.UsageViolation.lab_access_log_id,
            ).filter(models.UsageViolation.detected_at >= filters.start, models.UsageViolation.detected_at < until)
            for log in bounded(violations.order_by(models.UsageViolation.detected_at, models.UsageViolation.id)):
                row = sessions.get(log.lab_access_log_id)
                if row is None:
                    continue
                mark(_utc(log.detected_at).astimezone(LAB_TIMEZONE).date(), row)["violations"] += 1
                _append(violation_sheet, _common(row) + [log.id, _excel_time(log.detected_at), log.program_name,
                        log.reason, log.action_taken, log.detection_source])

        for sheet, count, heading in zip(sheets[1:],
                (len(aggregates), len(sessions), sum(day["violations"] for day in days.values())), headers):
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(heading))}{count + 1}"

        def unique(key):
            return len({getattr(row, key) for row in sessions.values() if getattr(row, key) is not None})

        summary.column_dimensions["A"].width = 50
        summary.column_dimensions["B"].width = 32
        for column in "CDEFG":
            summary.column_dimensions[column].width = 24
        summary.row_dimensions[2].height = 28
        metadata = [[], [labels["title"]], [], [labels["from"], filters.date_from], [labels["to"], filters.date_to],
                    [labels["generated"], _excel_time(now)], [labels["timezone"], "Asia/Bangkok (UTC+07:00)"],
                    [labels["lab"], filters.lab_id or labels["all"]], [labels["device"], filters.device_id or labels["all"]],
                    [labels["user"], filters.user_id or labels["all"]], [],
                    [labels["users"], unique("user_id")], [labels["devices"], unique("lab_device_id")],
                    [labels["sessions"], len(sessions)], [labels["session_seconds"], sum(d["session_seconds"] for d in days.values())],
                    [labels["usage_seconds"], sum(d["usage_seconds"] for d in days.values())],
                    [labels["violations"], sum(d["violations"] for d in days.values())], [labels["invalid"], invalid], []]
        for values in metadata:
            _append(summary, values)
        _append(summary, labels["daily"], header=True)
        for day, totals in days.items():
            _append(summary, [day, len(totals["users"]), len(totals["devices"]), len(totals["sessions"]),
                             totals["session_seconds"], totals["usage_seconds"], totals["violations"]])
        _append(summary, [])
        for offset, key in enumerate(("definition", "snapshot", "history")):
            summary.row_dimensions[len(metadata) + len(days) + 3 + offset].height = 90
            _append(summary, [labels["notes"], labels[key]])
        workbook.save(output)
        output.seek(0)
        return output
    except Exception:
        output.close()
        raise
    finally:
        # Close write-only XML writers, including aborted oversized exports.
        for sheet in sheets:
            if not sheet.closed:
                sheet.close()
            # openpyxl normally removes these XML spools during save(). An
            # aborted/oversized export must release them without waiting for
            # process exit. This writer hook is covered against the pinned version.
            writer = getattr(sheet, "_writer", None)
            if writer is not None:
                try:
                    writer.cleanup()
                except FileNotFoundError:
                    pass
        workbook.close()
